import csv
import io

from app import timesheet_exports
from app.models import Employee
from openpyxl import load_workbook
from test_periods_adjustments import correction, lock, prepared, snapshot
from test_shifts import schedule as schedule
from test_timesheets import add_events, get_report


def query(report, format="csv", **extra):
    return {
        "format": format,
        "source": "preview",
        "source_digest": report["source_digest"],
        "date_from": report["period"]["date_from"],
        "date_to": report["period"]["date_to"],
        "as_of": report["as_of"],
    } | extra


def test_preview_csv_and_xlsx_match_screen_keep_missing_blank_and_export_all_pages(schedule):
    browser, factory, headers, people, _, _ = prepared(schedule, missing=True)
    first = correction(browser, headers, people[0])
    assert first.status_code == 201
    # One missing second employee, so a missing value cannot silently become zero.
    _, _, _, _, create, assign = schedule
    shift = create()
    assert assign(shift, begin="2025-12-31", employees=[people[1]]).status_code == 201
    with factory() as db:
        db.get(Employee, people[0]).code = "000123"
        db.commit()
    report = get_report(browser, limit=1)
    assert len(report["rows"]) == 1 and report["has_more"]
    csv_response = browser.get("/api/admin/timesheets/export", params=query(report))
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    shifts = [row for row in rows if row["Loại dòng"] == "Ca làm"]
    assert len(shifts) == 2
    assert sorted(row["Giây trong ca"] for row in shifts) == ["", "28800.0"]
    assert all(row["Mã nguồn dữ liệu"] == report["source_digest"] for row in rows)
    excel = browser.get("/api/admin/timesheets/export", params=query(report, "xlsx"))
    assert excel.status_code == 200
    workbook = load_workbook(io.BytesIO(excel.content))
    assert workbook.sheetnames == ["Bảng công", "Ngoại lệ", "Thông tin", "Lượt tính công"]
    cells = list(workbook["Bảng công"].values)
    assert len(cells) == 3
    code_cells = [row[2] for row in workbook["Bảng công"].iter_rows(min_row=2)]
    assert any(cell.value == "000123" and cell.data_type == "s" for cell in code_cells)
    assert sorted(row[10] for row in cells[1:] if row[10] is not None) == [480]
    assert sum(row[10] is None for row in cells[1:]) == 1
    assert any(row[-1] == "Bổ sung bởi quản trị" for row in workbook["Lượt tính công"].values)
    assert excel.headers["x-report-source-digest"] == report["source_digest"]
    assert excel.headers["cache-control"] == "no-store"


def test_export_refuses_changed_data_and_snapshot_stays_original_after_profile_edit(schedule):
    browser, _, headers, people, _, _ = prepared(schedule)
    original = get_report(browser)
    assert (
        browser.patch(
            f"/api/admin/employees/{people[0]}", headers=headers, json={"name": "=TEST(1)"}
        ).status_code
        == 200
    )
    assert browser.get("/api/admin/timesheets/export", params=query(original)).status_code == 409
    current = get_report(browser)
    exported = browser.get("/api/admin/timesheets/export", params=query(current, "xlsx"))
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook["Bảng công"]["D2"].value == "=TEST(1)"
    assert workbook["Bảng công"]["D2"].data_type == "s"
    csv_response = browser.get("/api/admin/timesheets/export", params=query(current))
    assert "'=TEST(1)" in csv_response.content.decode("utf-8-sig")
    assert lock(browser, headers).status_code == 200
    frozen = snapshot(browser)
    assert (
        browser.patch(
            f"/api/admin/employees/{people[0]}", headers=headers, json={"name": "Tên sau chốt"}
        ).status_code
        == 200
    )
    params = {
        "format": "xlsx",
        "source": "snapshot",
        "source_digest": frozen["source_digest"],
        "month": "2025-12",
        "revision": 1,
    }
    old = browser.get("/api/admin/timesheets/export", params=params)
    assert old.status_code == 200
    assert load_workbook(io.BytesIO(old.content))["Bảng công"]["D2"].value == "=TEST(1)"
    assert "v1.xlsx" in old.headers["content-disposition"]
    assert original["source_digest"] != current["source_digest"]


def test_changed_adjustment_rejects_stale_export_and_orphan_rows_are_included(schedule):
    browser, _, headers, people, _, _ = prepared(schedule)
    old = get_report(browser)
    assert correction(browser, headers, people[0], occurred_at="2025-12-31T20:00:00+07:00").status_code == 201
    assert browser.get("/api/admin/timesheets/export", params=query(old)).status_code == 409
    current = get_report(browser)
    assert current["summary"]["unassigned_count"] == 1
    response = browser.get("/api/admin/timesheets/export", params=query(current))
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert sum(row["Loại dòng"] == "Lượt chưa gắn ca" for row in rows) == 1


def test_export_auth_bounds_capacity_and_invalid_reference(schedule):
    browser, _, _, _, _, _ = prepared(schedule)
    report = get_report(browser)
    assert browser.get("/api/admin/timesheets/export", params=query(report, "pdf")).status_code == 422
    assert (
        browser.get(
            "/api/admin/timesheets/export", params=query(report, source="snapshot", month="2025-12")
        ).status_code
        == 422
    )
    assert (
        browser.get("/api/admin/timesheets/export", params=query(report, source_digest="bad")).status_code
        == 422
    )
    assert (
        browser.get(
            "/api/admin/timesheets/export", params=query(report, as_of="2000-01-01T00:00:00Z")
        ).status_code
        == 422
    )
    for _ in range(2):
        assert timesheet_exports.slots.acquire(blocking=False)
    try:
        assert browser.get("/api/admin/timesheets/export", params=query(report)).status_code == 429
    finally:
        for _ in range(2):
            timesheet_exports.slots.release()
    browser.cookies.clear()
    assert browser.get("/api/admin/timesheets/export", params=query(report)).status_code == 401


def test_overnight_month_end_export_includes_next_day_adjustment_and_vietnamese_name(schedule):
    browser, factory, headers, people, create, assign = schedule
    shift = create({"name": "Ca đêm tiếng Việt", "start": "22:00", "end": "06:00"})
    assert assign(shift, begin="2025-12-31", employees=[people[0]]).status_code == 201
    add_events(factory, people[0], [("in", "2025-12-31T22:05:00+07:00")])
    assert correction(browser, headers, people[0], occurred_at="2026-01-01T05:55:00+07:00").status_code == 201
    report = get_report(browser)
    assert report["rows"][0]["worked_seconds"] == 470 * 60
    assert lock(browser, headers).status_code == 200
    frozen = snapshot(browser)
    params = {
        "source": "snapshot",
        "source_digest": frozen["source_digest"],
        "month": "2025-12",
        "revision": 1,
    }
    for format in ["csv", "xlsx"]:
        response = browser.get("/api/admin/timesheets/export", params=params | {"format": format})
        assert response.status_code == 200
        if format == "csv":
            rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
            row = next(row for row in rows if row["Loại dòng"] == "Ca làm")
            assert row["Ca làm"] == "Ca đêm tiếng Việt" and float(row["Phút trong ca"]) == 470
            assert row["Ngày ca"] == "2025-12-31"
        else:
            book = load_workbook(io.BytesIO(response.content))
            row = list(book["Bảng công"].values)[1]
            assert row[1] == "2025-12-31" and row[4] == "Ca đêm tiếng Việt" and row[10] == 470
            assert any(row[-1] == "Bổ sung bởi quản trị" for row in book["Lượt tính công"].values)
