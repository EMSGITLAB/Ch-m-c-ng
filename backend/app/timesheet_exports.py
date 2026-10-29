"""Local CSV/XLSX exports bound to the exact report displayed by the admin."""

import csv
import io
import re
import threading
from datetime import UTC, date, datetime, timedelta
from typing import Literal

from fastapi import APIRouter, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from pydantic import AwareDatetime
from sqlalchemy.exc import SQLAlchemyError
from starlette.concurrency import run_in_threadpool

from .auth import Admin, Db
from .periods import snapshot_report
from .timesheets import build_report

router = APIRouter(prefix="/api/admin/timesheets", tags=["timesheets"])
slots = threading.BoundedSemaphore(2)
STATUSES = {
    "complete": "Đã ghép đủ",
    "incomplete": "Thiếu dữ liệu",
    "needs_review": "Cần xem lại",
    "pending": "Đang chờ lượt",
    "in_progress": "Đang trong ca",
    "upcoming": "Chưa bắt đầu",
}
HEADERS = [
    "Loại dòng",
    "Ngày ca",
    "Mã nhân viên",
    "Họ tên",
    "Ca làm",
    "Phiên bản ca",
    "Trạng thái",
    "Vào đầu",
    "Ra cuối",
    "Giây trong ca",
    "Phút trong ca",
    "Phút dự kiến",
    "Phút muộn",
    "Phút sớm",
    "Phút ngoài ca",
    "Vấn đề",
    "Thời điểm lượt",
    "Chiều",
    "Nguồn lượt",
    "Mã lượt",
    "Mã phân ca",
]
META_HEADERS = [
    "Từ ngày",
    "Đến ngày",
    "Múi giờ",
    "Thời điểm tính",
    "Loại báo cáo",
    "Phiên bản chốt",
    "Mã nguồn dữ liệu",
    "Số lượt thử chức năng",
]


def export_rows(report):
    rows = []
    for row in report["rows"]:
        rows.append(
            [
                "Ca làm",
                row["work_date"],
                row["employee_code"],
                row["employee_name"],
                row["shift"]["name"],
                row["shift_revision"],
                STATUSES.get(row["status"], row["status"]),
                row["first_in"],
                row["last_out"],
                row["worked_seconds"],
                None if row["worked_seconds"] is None else row["worked_seconds"] / 60,
                row["expected_seconds"] / 60,
                None if row["late_seconds"] is None else row["late_seconds"] / 60,
                None if row["early_seconds"] is None else row["early_seconds"] / 60,
                row["outside_shift_seconds"] / 60,
                "; ".join(row["exceptions"]),
                None,
                None,
                None,
                None,
                row["assignment_id"],
            ]
        )
    for group, label in [("unassigned_events", "Lượt chưa gắn ca"), ("ambiguous_events", "Lượt mơ hồ")]:
        for row in report[group]:
            rows.append(
                [
                    label,
                    None,
                    row["employee_code"],
                    row["employee_name"],
                    None,
                    None,
                    label,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "; ".join(row.get("candidate_assignments", [])),
                    row["at"],
                    "Vào" if row["direction"] == "in" else "Ra",
                    None,
                    row["event_id"],
                    None,
                ]
            )
    return rows


def csv_text(value):
    if not isinstance(value, str):
        return value
    # CSV consumers may execute leading formulas even inside quoted strings.
    return (
        "'" + value
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n"))
        else value
    )


def xml_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda match: f"\\u{ord(match[0]):04x}", value)


def serialize_export(report, format):
    rows = export_rows(report)
    state = (
        "Bản chốt hiện hành"
        if report["finalized"] and report.get("is_current")
        else "Bản chốt lưu đối chiếu"
        if report["finalized"]
        else "Bản xem trước"
    )
    metadata = [
        report["period"]["date_from"],
        report["period"]["date_to"],
        report["period"]["timezone"],
        report["as_of"],
        state,
        report.get("snapshot_revision"),
        report["source_digest"],
        report.get("functional_trial_event_count", 0),
    ]
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(HEADERS + META_HEADERS)
        writer.writerow(["Thông tin báo cáo"] + [None] * (len(HEADERS) - 1) + metadata)
        for row in rows:
            writer.writerow([csv_text(value) for value in row + metadata])
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    book = Workbook(write_only=True)

    def sheet(name, headers, values):
        target = book.create_sheet(name)
        target.freeze_panes = "A2"
        target.sheet_view.showGridLines = False
        for index, title in enumerate(headers, 1):
            from openpyxl.utils import get_column_letter

            target.column_dimensions[get_column_letter(index)].width = min(42, max(18, len(title) + 3))

        def append(values, heading=False):
            cells = []
            for value in values:
                cell = WriteOnlyCell(target, xml_text(value) if isinstance(value, str) else value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Preserve codes, and never interpret names/reasons as formulas.
                elif isinstance(value, (float, int)):
                    cell.number_format = "0.00######"
                if heading:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="42613D")
                cells.append(cell)
            target.append(cells)

        append(headers, True)
        for value in values:
            append(value)
        return target

    sheet("Bảng công", HEADERS, [row for row in rows if row[0] == "Ca làm"])
    sheet("Ngoại lệ", HEADERS, [row for row in rows if row[0] != "Ca làm"])
    sheet(
        "Thông tin",
        ["Nội dung", "Giá trị"],
        list(zip(META_HEADERS, metadata, strict=True))
        + [
            ["Phiên bản tính công", report["calculation_version"]],
            ["Mã bản chốt", report.get("snapshot_id")],
            ["Lý do chốt", report.get("snapshot_reason")],
            ["Tổng số ca", report["summary"]["row_count"]],
            ["Tổng giây đã xác định", report["summary"]["worked_seconds"]],
            ["Lượt chưa gắn ca", report["summary"]["unassigned_count"]],
            ["Lượt mơ hồ", report["summary"]["ambiguous_count"]],
            [
                "Lưu ý",
                "Ca thiếu/mơ hồ để trống công; giờ ngoài ca chưa tự tính tăng ca. Thời điểm có độ lệch UTC; múi giờ hiển thị ở thông tin báo cáo.",
            ],
        ],
    )
    sheet(
        "Lượt tính công",
        ["Ngày ca", "Mã nhân viên", "Họ tên", "Mã phân ca", "Mã lượt", "Chiều", "Thời điểm", "Nguồn"],
        [
            [
                row["work_date"],
                row["employee_code"],
                row["employee_name"],
                row["assignment_id"],
                item["id"],
                "Vào" if item["direction"] == "in" else "Ra",
                item["at"],
                "Bổ sung bởi quản trị" if item.get("source") == "manual" else "Camera",
            ]
            for row in report["rows"]
            for item in row["events"]
        ],
    )
    output = io.BytesIO()
    book.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_export(
    db,
    format,
    source,
    source_digest,
    date_from=None,
    date_to=None,
    month=None,
    revision=None,
    employee_id=None,
    as_of=None,
):
    if source == "snapshot":
        if not month or revision is None:
            raise HTTPException(422, "Chọn rõ tháng và phiên bản bản chốt cần xuất.")
        report = snapshot_report(db, month, revision, employee_id, unpaginated=True)
    else:
        now = datetime.now(UTC)
        if (
            date_from is None
            or date_to is None
            or as_of is None
            or not now - timedelta(days=1) <= as_of <= now
        ):
            raise HTTPException(
                422, "Cần ngày và thời điểm của bản xem trước trong 24 giờ qua. Hãy tính lại bảng công."
            )
        report = build_report(db, date_from, date_to, employee_id=employee_id, as_of=as_of, unpaginated=True)
    if report["source_digest"] != source_digest:
        raise HTTPException(409, "Dữ liệu đã thay đổi. Tải lại bảng công trước khi xuất báo cáo.")
    body, mime = serialize_export(report, format)
    version = f"v{report['snapshot_revision']}" if report["finalized"] else "preview"
    filename = f"bang-cong_{report['period']['date_from']}_{report['period']['date_to']}_{version}.{format}"
    return Response(
        content=body,
        media_type=mime,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Report-Source-Digest": source_digest,
        },
    )


@router.get("/export")
async def download(
    db: Db,
    user: Admin,
    format: Literal["csv", "xlsx"],
    source: Literal["preview", "snapshot"],
    source_digest: str = Query(pattern=r"^[0-9a-f]{64}$"),
    date_from: date | None = None,
    date_to: date | None = None,
    month: str | None = None,
    revision: int | None = None,
    employee_id: str | None = None,
    as_of: AwareDatetime | None = None,
):
    if not slots.acquire(blocking=False):
        raise HTTPException(429, "Đang xuất báo cáo khác. Vui lòng thử lại.")
    try:
        return await run_in_threadpool(
            build_export,
            db,
            format,
            source,
            source_digest,
            date_from,
            date_to,
            month,
            revision,
            employee_id,
            as_of,
        )
    except (ValueError, OverflowError) as error:
        raise HTTPException(422, "Dữ liệu hoặc khoảng thời gian chưa hợp lệ; chưa xuất báo cáo.") from error
    except SQLAlchemyError as error:
        raise HTTPException(503, "Chưa đọc đủ dữ liệu; chưa xuất báo cáo. Vui lòng thử lại.") from error
    finally:
        slots.release()
